import argparse,json,sys
from pathlib import Path
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

def money(x):
 try:d=Decimal(str(x))
 except InvalidOperation:raise ValueError('Invalid amount')
 if not d.is_finite() or d<0 or d>Decimal('10000000') or d.as_tuple().exponent < -2:raise ValueError('Amount must be nonnegative, <=10000000 and have at most two decimal places')
 return d

def text(x,n=100):
 if not isinstance(x,str) or not x.strip() or len(x)>n or any(ord(c)<32 for c in x):raise ValueError('Invalid text')
 return x.strip()

def validate(data):
 if not isinstance(data,dict):raise ValueError('Expected object')
 out={'number':text(data.get('number'),50),'customer':text(data.get('customer')),'paid':money(data.get('paid',0)),'lines':[]}
 lines=data.get('lines')
 if not isinstance(lines,list) or not 1<=len(lines)<=50:raise ValueError('Need 1–50 lines')
 for l in lines:
  if not isinstance(l,dict):raise ValueError('Line must be object')
  q=l.get('quantity');tax=money(l.get('tax_percent',0));price=money(l.get('unit_price'))
  if type(q)!=int or not 1<=q<=10000 or tax>100:raise ValueError('Invalid quantity or tax')
  out['lines'].append({'description':text(l.get('description')),'quantity':q,'unit_price':price,'tax_percent':tax})
 return out

def totals(data):
 d=validate(data);net=Decimal(0);tax=Decimal(0)
 for l in d['lines']:
  amount=(l['unit_price']*l['quantity']).quantize(Decimal('.01'),rounding=ROUND_HALF_UP);net+=amount;tax+=(amount*l['tax_percent']/100).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
 return dict(net=float(net),tax=float(tax),total=float(net+tax),balance=float(net+tax-d['paid']))

def literal(cell,value):cell.value=value;cell.data_type='s'

def build(data,path):
 d=validate(data);p=Path(path)
 if p.exists():raise ValueError('Output already exists')
 w=Workbook();s=w.active;s.title='Invoice';w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
 s.sheet_view.showGridLines=False;s.merge_cells('A1:G1');s['A1']='INVOICE / QUOTATION';s['A1'].font=Font(size=24,bold=True,color='173C40')
 s.merge_cells('A2:G2');s['A2']='SYNTHETIC SAMPLE • Portfolio demo • Currency INR';s['A2'].font=Font(color='627C7D')
 s['A4']='Number';literal(s['B4'],d['number']);s['A5']='Customer';s.merge_cells('B5:G5');literal(s['B5'],d['customer'])
 headers=['Description','Qty','Unit INR','Tax %','Net INR','Tax INR','Total INR']
 for i,h in enumerate(headers,1):c=s.cell(7,i,h);c.fill=PatternFill('solid',fgColor='173C40');c.font=Font(color='FFFFFF',bold=True)
 for r,l in enumerate(d['lines'],8):
  literal(s.cell(r,1),l['description']);s.cell(r,2,l['quantity']);s.cell(r,3,float(l['unit_price']));s.cell(r,4,float(l['tax_percent']))
  s.cell(r,5,f'=ROUND(B{r}*C{r},2)');s.cell(r,6,f'=ROUND(E{r}*D{r}/100,2)');s.cell(r,7,f'=E{r}+F{r}')
  for col in range(1,8):
   c=s.cell(r,col);c.fill=PatternFill('solid',fgColor=('EAF3FD' if col<5 else 'F1F6F4'));c.alignment=Alignment(vertical='center',wrap_text=True)
   if col>=3:c.number_format='#,##0.00;[Red]-#,##0.00'
  s.row_dimensions[r].height=32
 end=7+len(d['lines']);r=end+2
 for row,label,formula in ((r,'Subtotal',f'=SUM(E8:E{end})'),(r+1,'Tax',f'=SUM(F8:F{end})'),(r+2,'Grand total',f'=SUM(G8:G{end})'),(r+3,'Paid',float(d['paid'])),(r+4,'Balance',f'=G{r+2}-G{r+3}')):
  s.cell(row,6,label);s.cell(row,7,formula);s.cell(row,7).number_format='#,##0.00;[Red]-#,##0.00';s.cell(row,6).font=Font(bold=True);s.cell(row,7).font=Font(bold=True,color='173C40')
 s.cell(r+3,7).fill=PatternFill('solid',fgColor='EAF3FD')
 s.cell(r+6,1,'Blue cells are editable inputs. Totals recalculate when opened. Negative balance means overpayment.')
 for col,width in {'A':38,'B':9,'C':14,'D':11,'E':15,'F':15,'G':16}.items():s.column_dimensions[col].width=width
 s.freeze_panes='B8';s.print_options.horizontalCentered=True;s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1;s.print_area=f'A1:G{r+7}'
 for cells,lo,hi,typ in ((f'B8:B{end}',1,10000,'whole'),(f'C8:C{end}',0,10000000,'decimal'),(f'D8:D{end}',0,100,'decimal')):
  dv=DataValidation(type=typ,operator='between',formula1=lo,formula2=hi);dv.showErrorMessage=True;dv.error='Enter a value within the permitted range';s.add_data_validation(dv);dv.add(cells)
 s['I1']='Validation target (generator)';s['I2']='Synthetic data only';s.column_dimensions['I'].hidden=True
 w.save(p);return {'sheet':'Invoice','subtotal_cell':f'G{r}','tax_cell':f'G{r+1}','total_cell':f'G{r+2}','balance_cell':f'G{r+4}',**totals(data)}

def main():
 ap=argparse.ArgumentParser(description='Generate a formula-driven invoice workbook');ap.add_argument('json');ap.add_argument('output');a=ap.parse_args()
 try:print(json.dumps(build(json.loads(Path(a.json).read_text()),a.output),indent=2));return 0
 except (ValueError,OSError,TypeError) as e:print(f'Error: {e}',file=sys.stderr);return 2
if __name__=='__main__':sys.exit(main())
